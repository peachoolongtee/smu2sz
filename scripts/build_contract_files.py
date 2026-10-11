"""Build the .docx and text-PDF copies of the demo contracts (fixtures/contracts/contracts.yaml).

    uv run --project services/api python scripts/build_contract_files.py

The .md file is the canonical text. The copies exercise the upload path (app/ingest): Word via python-docx,
PDF via a minimal writer below (Helvetica, WinAnsi, one text line per Tj), so no PDF library is needed to make
them. The PDF is reproducible byte for byte; the .docx is a zip with entry timestamps, so only its text is.
"""

from __future__ import annotations

import sys
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
CONTRACTS = ROOT / "fixtures" / "contracts"


def blocks(text: str) -> list[str]:
    """Paragraphs of the .md text: blank-line separated; a block's own lines (list items) stay separate."""
    out: list[str] = []
    for chunk in text.strip("\n").split("\n\n"):
        out += [ln for ln in chunk.split("\n") if ln.strip()]
    return out


def is_heading(line: str) -> bool:
    letters = [c for c in line if c.isalpha()]
    return bool(letters) and all(c.isupper() for c in letters)


# ---------------------------------------------------------------- Word


def write_docx(text: str, path: Path) -> None:
    import docx

    document = docx.Document()
    core = document.core_properties  # fixed metadata: the same input gives the same document text
    core.author, core.title, core.comments = "TermsLedger demo", path.stem, "Synthetic demonstration document"
    for i, line in enumerate(blocks(text)):
        if i == 0:
            document.add_heading(line, level=1)
        elif is_heading(line):
            document.add_heading(line, level=2)
        else:
            document.add_paragraph(line)
    document.save(str(path))


# ---------------------------------------------------------------- PDF

# Helvetica advance widths (1/1000 em) for printable ASCII 32..126, from the standard AFM.
_W = [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278, 556, 556, 556, 556, 556, 556,
      556, 556, 556, 556, 278, 278, 584, 584, 584, 556, 1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667,
      556, 833, 722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556, 333, 556,
      556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556, 556, 556, 333, 500, 278, 556, 500, 722,
      500, 500, 500, 334, 260, 334, 584]
FONT_SIZE = 10
LEADING = 14
PAGE_W, PAGE_H, MARGIN = 595, 842, 56
MAX_W = (PAGE_W - 2 * MARGIN) * 1000 / FONT_SIZE  # in 1/1000 em units


def _width(s: str) -> int:
    return sum(_W[ord(c) - 32] if 32 <= ord(c) <= 126 else 556 for c in s)


def wrap(line: str) -> list[str]:
    words, out, cur = line.split(" "), [], ""
    for w in words:
        cand = f"{cur} {w}" if cur else w
        if _width(cand) <= MAX_W or not cur:
            cur = cand
        else:
            out.append(cur)
            cur = w
    out.append(cur)
    return out


def _escape(s: str) -> bytes:
    raw = s.encode("cp1252", errors="replace")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def write_pdf(text: str, path: Path) -> None:
    lines: list[str | None] = []  # None is a blank line between paragraphs
    for b in blocks(text):
        lines += wrap(b)
        lines.append(None)
    per_page = (PAGE_H - 2 * MARGIN) // LEADING
    pages = [lines[i:i + per_page] for i in range(0, len(lines), per_page)]
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")  # filled once the page ids are known
    pages_id = add(b"")
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    page_ids = []
    for page in pages:
        ops = [f"BT /F1 {FONT_SIZE} Tf {LEADING} TL {MARGIN} {PAGE_H - MARGIN} Td".encode()]
        for ln in page:
            ops.append(b"T*" if ln is None else b"(" + _escape(ln) + b") Tj T*")
        ops.append(b"ET")
        stream = b"\n".join(ops)
        contents = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(add(
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 %d %d] /Resources << /Font << /F1 %d 0 R >> >> "
            b"/Contents %d 0 R >>" % (pages_id, PAGE_W, PAGE_H, font, contents)))
    objects[catalog - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id
    kids = b" ".join(b"%d 0 R" % i for i in page_ids)
    objects[pages_id - 1] = b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(page_ids)

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % n + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref)
    path.write_bytes(bytes(out))


def main() -> int:
    registry = yaml.safe_load((CONTRACTS / "contracts.yaml").read_text(encoding="utf-8"))["contracts"]
    for c in registry:
        src = CONTRACTS / c["file"]
        text = src.read_text(encoding="utf-8")
        for fmt in c.get("formats", []):
            dest = src.with_suffix(f".{fmt}")
            (write_docx if fmt == "docx" else write_pdf)(text, dest)
            print(f"wrote {dest.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
