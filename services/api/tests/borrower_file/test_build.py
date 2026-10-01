"""Borrower file: built from verified state only, every figure labelled, nothing private leaks."""

from __future__ import annotations

import io

import pytest

from app.borrower_file import (
    TITLE_EN,
    TITLE_ZH,
    BorrowerFileError,
    Evidence,
    build,
    customer_ref,
    pdf_renderer,
    render_docx,
    render_html,
    render_pdf,
)
from app.counterparties.profiles import load_buyer_profiles
from app.domain.fixtures import load_scenario
from app.domain.models import Actor, DrawFacility, FinancialState
from app.egress import AllowedValue, Category, ProtectedSet, ValueKind, assert_no_leak
from app.mandate import demo_owner_public_key_hex, load_signed_mandate
from app.operations import Executor
from app.persistence import InMemoryLedgerStore
from tests.egress.helpers import reservation_private_values, state_with_buyer_b

REQUEST = {"facility_id": "FAC-GGL", "amount_fen": 36_000_000, "use_of_proceeds": "Working capital for payroll"}


def _store(state: FinancialState | None = None) -> tuple[InMemoryLedgerStore, Executor]:
    store = InMemoryLedgerStore()
    store.put_state(state or load_scenario()[0])
    store.put_mandate(load_signed_mandate())
    return store, Executor(store, trusted_owner_key_hex=demo_owner_public_key_hex(), clock=lambda: 0)


def _draw(ex: Executor, amount: int = 15_000_000):
    return ex.execute(DrawFacility(op_id="OP-DRAW-1", idempotency_key="draw-1", mandate_version=1,
                                   requested_by=Actor.AGENT, day=0, facility_id="FAC-LINE", amount_fen=amount))


@pytest.fixture
def after_draw():
    store, ex = _store()
    result = _draw(ex)
    assert result.status.value == "EXECUTED"
    return store, result, build(store, day=0, **REQUEST)


def _figures(bf):
    return [f for s in bf.sections for r in s.rows for f in r.figures]


def test_file_after_a_line_draw_lists_the_draw_with_its_hash(after_draw):
    store, result, bf = after_draw
    ledger = next(s for s in bf.sections if s.key == "commitments")
    row = next(r for r in ledger.rows if r.label.startswith("Draw:"))
    assert "FAC-LINE" in row.label and f"seq {result.seq}" in row.label
    assert row.figures[0].fen == 15_000_000 and row.figures[0].evidence == Evidence.OBSERVED
    assert any(result.entry_hash in f.text for f in row.figures)
    assert bf.provenance.chain_head_hash == result.entry_hash == store.operations()[-1].entry_hash
    assert bf.provenance.chain_valid and bf.provenance.state_version == store.get_state().state_version
    assert bf.provenance.mandate_version == 1 and bf.provenance.generated_day == 0
    assert result.entry_hash in render_html(bf)


def test_every_figure_carries_an_evidence_label(after_draw):
    _, _, bf = after_draw
    figs = _figures(bf)
    assert figs and all(isinstance(f.evidence, Evidence) for f in figs)
    assert {f.evidence for f in figs} == set(Evidence)
    # the HTML shows one badge per figure
    assert render_html(bf).count('class="badge ') == len(figs)
    by = {r.label: r for s in bf.sections for r in s.rows}
    assert by["Amount requested"].figures[0].evidence == Evidence.OWNER_ESTIMATE
    assert by["Use of proceeds"].figures[0].evidence == Evidence.OWNER_ESTIMATE
    assert by["Cash on hand"].figures[0].evidence == Evidence.OBSERVED
    assert by["Cash on hand"].figures[0].fen == 55_000_000  # ¥400,000 + the ¥150,000 draw


def test_sections_cover_spec_7_8(after_draw):
    _, _, bf = after_draw
    keys = [s.key for s in bf.sections]
    assert keys == ["summary", "request", "forecast_13w", "trajectory_12m", "covenant_headroom", "covenant_tests",
                    "commitments", "proceeds", "evidence", "provenance"]
    sec = {s.key: s for s in bf.sections}
    assert len(sec["forecast_13w"].rows) == 13
    assert all(f.evidence == Evidence.PROJECTED for r in sec["forecast_13w"].rows for f in r.figures)
    assert len(sec["trajectory_12m"].rows) == 12 and len(sec["covenant_headroom"].rows) == 12
    # week 2 holds the day-12 payroll; the pro-forma facility lands on day 21 (21-day lead time)
    wk2, wk3 = sec["forecast_13w"].rows[1], sec["forecast_13w"].rows[2]
    assert wk2.figures[1].fen == -34_000_000
    assert wk3.figures[4].fen - wk3.figures[2].fen == 36_000_000
    # leverage rises after the pro-forma draw; DSCR headroom against FAC-LINE's 1.25x. The pro-forma ¥360,000
    # guaranteed loan is a bullet term loan, serviced on a level basis (ADR-0008): DSCR 1.18x, headroom -0.07x
    m1 = sec["covenant_headroom"].rows[0]
    assert m1.figures[0].text == "2.00×" and m1.figures[2].text == "−0.07×"
    assert m1.figures[3].text == "1.53×" and m1.figures[4].text == "1.86×"
    assert bf.status == "PREPARED_NOT_SENT" and bf.lender_mocked is True


def test_html_is_self_contained_bilingual_and_neutral(after_draw):
    _, _, bf = after_draw
    html = render_html(bf)
    assert html.startswith("<!doctype html>") and TITLE_EN in html and TITLE_ZH in html
    assert "mocked" in html and "prepared, not sent" in html.lower()
    for zh in ("经营概况", "13周现金流预测", "承诺台账", "数据来源"):
        assert zh in html
    assert "<script" not in html and "http://" not in html and "https://" not in html
    assert "webank" not in html.lower() and "微众" not in html and "微业贷" not in html


def test_html_leaks_no_other_buyer_values_to_the_lender():
    state, mandate = state_with_buyer_b()  # a second buyer with its own invoice and 60-day terms
    store, ex = _store(state)
    _draw(ex)
    bf = build(store, day=0, **REQUEST)
    html = render_html(bf)
    extra = [pv for p in load_buyer_profiles().values() for pv in reservation_private_values(p)]
    ps = ProtectedSet.build(store.get_state(), mandate, extra_private=extra)
    # What the owner deliberately discloses to a lender: its own balances and facilities, its receivables
    # and booked concessions as amounts, and covenant ratios. Never: counterparty names or ids, a buyer's
    # terms, buyer reservation values, the mandate's money limits or canaries.
    disclosed = [AllowedValue(v.kind, v.value) for v in ps.values
                 if v.category in (Category.BALANCE, Category.FACILITY, Category.CONCESSION)
                 or (v.category == Category.COUNTERPARTY_TERMS and v.kind == ValueKind.MONEY)
                 or (v.category == Category.MANDATE and v.kind == ValueKind.RATIO)]
    assert_no_leak(html, ps, "BANK-2", allow=disclosed)
    for cp in state.counterparties:
        if cp.role != "OWNER":
            assert cp.name not in html and cp.id not in html
    assert customer_ref("BUYER-1") in html


def test_docx_opens_with_python_docx(after_draw):
    from docx import Document

    _, result, bf = after_draw
    doc = Document(io.BytesIO(render_docx(bf)))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == TITLE_EN and TITLE_ZH in texts
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert any(result.entry_hash in c for c in cells)
    assert len(doc.tables) == len(bf.sections)


def test_pdf_is_skipped_without_a_local_renderer(after_draw):
    _, _, bf = after_draw
    if pdf_renderer() is None:
        assert render_pdf(bf) is None
    else:  # pragma: no cover - only where WeasyPrint happens to be installed
        assert render_pdf(bf)[:4] == b"%PDF"


def test_file_id_is_deterministic_and_tracks_state(after_draw):
    store, _, bf = after_draw
    assert build(store, day=0, **REQUEST).file_id == bf.file_id
    assert build(store, day=0, **{**REQUEST, "amount_fen": 20_000_000}).file_id != bf.file_id


def test_owner_estimates_are_shown_but_not_projected():
    store, _ = _store()
    plain = build(store, day=0, **REQUEST)
    est = build(store, day=0, ebitda_uplift_fen=1_000_000, **REQUEST)
    rows = {r.label: r for s in est.sections for r in s.rows}
    uplift = rows["Expected monthly EBITDA uplift (not used in any projection)"].figures[0]
    assert uplift.evidence == Evidence.OWNER_ESTIMATE and uplift.fen == 1_000_000
    def traj(bf):
        return next(s for s in bf.sections if s.key == "trajectory_12m").rows

    assert traj(plain) == traj(est)


@pytest.mark.parametrize("kw, msg", [
    ({"facility_id": "FAC-NOPE"}, "owner-approved"),
    ({"facility_id": "FAC-IF"}, "not a loan"),
    ({"amount_fen": 0}, "within"),
    ({"amount_fen": 60_000_000}, "within"),
    ({"use_of_proceeds": "   "}, "use of proceeds"),
])
def test_bad_requests_are_refused(kw, msg):
    store, _ = _store()
    with pytest.raises(BorrowerFileError, match=msg):
        build(store, day=0, **{**REQUEST, **kw})


def test_accepted_concessions_carry_terms_hash_and_redacted_source_spans():
    from app.commitments import terms_hash
    from app.domain.models import Commitment, CommitmentStatus
    from tests.verifier.conftest import compiled

    obligations = compiled("EX01_PACKAGE").obligations
    c = Commitment(id="C-1", offer_id="OFF-1", offer_version=1, counterparty_id="BUYER-1", verdict_id="V-1",
                   terms_hash=terms_hash(obligations), obligations=obligations, status=CommitmentStatus.ACTIVE,
                   nonce="n-1", expires_day=10, mandate_version=1, forward_value_fen=400_000)
    store, _ = _store()
    store.put_commitment(c)
    bf = build(store, day=0, **REQUEST)
    ledger = next(s for s in bf.sections if s.key == "commitments")
    row = next(r for r in ledger.rows if r.label.startswith("Concession C-1"))
    assert customer_ref("BUYER-1") in row.label and "BUYER-1" not in row.label
    assert f"terms hash {c.terms_hash}" in [f.text for f in row.figures]
    spans = [f for f in row.figures if f.source == "source span"]
    assert len(spans) == sum(len(o.spans) for o in obligations)
    assert all(f.evidence == Evidence.OBSERVED for f in row.figures)
    assert any(o.spans[0].text in f.text for o in obligations for f in spans)
