"""Text of an uploaded document: plain text or Markdown, Word (.docx) and PDF with a text layer.

The text is the untrusted document the compiler quotes from, so it must be byte-stable: the same file always
gives the same text, because quotes are matched verbatim and cache keys hash it. Every path ends in
`normalise()`: LF line endings, NFC, no control characters, no trailing spaces, at most one blank line in a row.
Paragraphs are separated by a blank line and so are PDF pages, so a quote stays readable.

Libraries: python-docx (MIT) and pypdf (BSD). No PyMuPDF (AGPL) and no OCR: a PDF without a text layer is
refused with a message that says what to upload instead.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from pathlib import PurePath

MAX_BYTES = 5 * 1024 * 1024  # an upload larger than this is refused before it is parsed
MAX_CHARS = 400_000  # about 120 pages of contract text
MAX_UNZIPPED_BYTES = 20 * 1024 * 1024  # a Word file whose parts add up to more than this once unzipped is refused unopened
MAX_PDF_PAGES = 60  # a PDF with more pages than this is refused before any page is read
MIN_CHARS_PER_PAGE = 20  # a PDF averaging fewer letters and digits per page than this has no text layer

TEXT_SUFFIXES = (".txt", ".md", ".markdown", ".text")
SUPPORTED = (*TEXT_SUFFIXES, ".docx", ".pdf")


class IngestError(ValueError):
    """The document cannot be read. `code` is a stable reason for the API; the message is for the owner."""

    code = "UNREADABLE_DOCUMENT"
    message_zh: str | None = None  # the same sentence for a Chinese page, where a limit has one

    def __init__(self, message: str, *, message_zh: str | None = None) -> None:
        super().__init__(message)
        if message_zh is not None:
            self.message_zh = message_zh


class UnsupportedFileType(IngestError):
    code = "UNSUPPORTED_FILE_TYPE"


class ScannedPdf(IngestError):
    code = "SCANNED_PDF"


class DocumentTooLarge(IngestError):
    code = "DOCUMENT_TOO_LARGE"


class EmptyDocument(IngestError):
    code = "EMPTY_DOCUMENT"


SCANNED_MESSAGE = "this PDF is a scan; upload a text PDF or Word file"


def extract_text(filename: str, data: bytes) -> str:
    """The normalised text of one document. Raises an `IngestError` subclass when it cannot be read."""
    if len(data) > MAX_BYTES:
        raise DocumentTooLarge(f"the file is larger than {MAX_BYTES // (1024 * 1024)} MB",
                               message_zh=f"文件超过 {MAX_BYTES // (1024 * 1024)} MB")
    suffix = PurePath(filename or "").suffix.lower()
    if suffix in TEXT_SUFFIXES:
        text = _decode(data)
    elif suffix == ".docx":
        text = _docx(data)
    elif suffix == ".pdf":
        text = _pdf(data)
    elif suffix == ".doc":
        raise UnsupportedFileType("old Word (.doc) files are not read; save it as .docx or as a text PDF")
    else:
        raise UnsupportedFileType(f"unsupported file type {suffix or '(none)'}; upload .docx, a text PDF, .txt or .md")
    text = normalise(text)
    if not text.strip():
        raise EmptyDocument("the document has no text")
    if len(text) > MAX_CHARS:
        raise DocumentTooLarge(f"the document is longer than {MAX_CHARS:,} characters",
                               message_zh=f"文档超过 {MAX_CHARS:,} 个字符")
    return text


# ---------------------------------------------------------------- normalisation

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0e-\x1f\x7f]")
_BLANK_RUN = re.compile(r"\n{3,}")


def normalise(text: str) -> str:
    """LF endings, NFC, no BOM or control characters, no trailing spaces, at most one blank line in a row, no
    leading or trailing blank lines. Idempotent: normalise(normalise(t)) == normalise(t)."""
    t = text.replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n").replace(" ", "\n").replace(" ", "\n")
    t = unicodedata.normalize("NFC", t).lstrip("﻿")
    t = _CONTROL.sub("", t)
    t = "\n".join(line.rstrip() for line in t.split("\n"))
    t = _BLANK_RUN.sub("\n\n", t)
    return t.strip("\n")


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise IngestError("the text file is not UTF-8; save it as UTF-8 and upload it again") from e


# ---------------------------------------------------------------- Word


def _docx(data: bytes) -> str:
    if not data.startswith(b"PK"):
        raise IngestError("the file is not a Word .docx document")
    _check_unzipped_size(data)
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 - a corrupt or mislabelled file of any kind
        raise IngestError("the Word document could not be opened") from e
    blocks: list[str] = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document).text
            if text.strip():
                blocks.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _check_unzipped_size(data: bytes) -> None:
    """Refuse a Word file whose parts add up to more than `MAX_UNZIPPED_BYTES` once unzipped, before python-docx loads any
    of it. The sizes are the archive's declared ones, and reading a part never returns more than its declared size."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as e:
        raise IngestError("the Word document could not be opened") from e
    if total > MAX_UNZIPPED_BYTES:
        limit = MAX_UNZIPPED_BYTES // (1024 * 1024)
        raise DocumentTooLarge(f"the Word file is larger than {limit} MB once unzipped; upload a smaller copy",
                               message_zh=f"这份 Word 文件解压后超过 {limit} MB，请上传较小的版本")


# ---------------------------------------------------------------- PDF


def _pdf(data: bytes) -> str:
    if not data.lstrip()[:5].startswith(b"%PDF"):
        raise IngestError("the file is not a PDF")
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestError("the PDF is password-protected; upload an unprotected copy")
        n_pages = len(reader.pages)
        if n_pages > MAX_PDF_PAGES:
            raise DocumentTooLarge(f"the PDF has {n_pages} pages; the limit is {MAX_PDF_PAGES}",
                                   message_zh=f"这份 PDF 有 {n_pages} 页，上限是 {MAX_PDF_PAGES} 页")
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001 - one unreadable page is an empty page
                pages.append("")
    except IngestError:
        raise
    except Exception as e:  # noqa: BLE001
        raise IngestError("the PDF could not be opened") from e
    letters = sum(1 for p in pages for ch in p if ch.isalnum())
    if not pages or letters < MIN_CHARS_PER_PAGE * len(pages):
        raise ScannedPdf(SCANNED_MESSAGE)
    return "\n\n".join(unwrap(p) for p in pages if p.strip())


# A line that starts a new block: a clause number ("4.2 ", "12. ", never a bare "2026 " or "12 "), a list item
# ("(a) ", "(iv) "), a schedule or article heading, a signature line or a bullet.
_BLOCK_START = re.compile(
    r"^\s*(?:\d{1,2}(?:\.\d{1,2})+\.?\s|\d{1,2}\.\s|\((?:[a-z]|[ivx]{1,4}|\d{1,2})\)\s"
    r"|(?:SCHEDULE|Schedule|ARTICLE|Article|ANNEX|Annex|APPENDIX|Appendix)\b|Signed\b|[-•·]\s)")
_SHORT_LINE = 0.8  # a line shorter than this share of the page's longest line ends its paragraph


def _heading(line: str) -> bool:
    """An all-capitals heading ("5. FINANCIAL COVENANTS"); a short capitalised fragment ("RMB 150,000.") is not."""
    letters = [c for c in line if c.isalpha()]
    return len(letters) >= 6 and all(c.isupper() for c in letters)


_TERMINAL = (".", ":", ";", "!", "?", "。", "；", "：")


def unwrap(page: str) -> str:
    """Re-join the visual line wraps of one PDF page into paragraphs, so a sentence is never cut by a wrap.

    A line joins the one before it unless it starts a block itself or is a heading, and the line before is not a
    heading, when either that line is nearly full width (at least 80% of the page's longest line) or it plainly
    continues (no terminal punctuation, and the new line starts in lower case). A word hyphenated at the wrap is
    re-joined without the hyphen only when the next line starts in lower case. Paragraphs are separated by a
    blank line."""
    lines = [ln.rstrip() for ln in page.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    width = max((len(ln) for ln in lines if ln.strip()), default=0)
    paragraphs: list[str] = []
    cur = ""
    prev = ""
    for line in lines:
        if not line.strip():
            if cur:
                paragraphs.append(cur)
            cur, prev = "", ""
            continue
        stripped = line.strip()
        continues = len(prev) >= _SHORT_LINE * width or (
            not prev.rstrip().endswith(_TERMINAL) and stripped[:1].islower())
        joins = (bool(cur) and continues and not _heading(prev)
                 and not _BLOCK_START.match(line) and not _heading(line))
        if joins:
            if cur.endswith("-") and stripped[:1].islower():
                cur = cur[:-1] + stripped
            else:
                cur = f"{cur} {stripped}"
        else:
            if cur:
                paragraphs.append(cur)
            cur = stripped
        prev = line
    if cur:
        paragraphs.append(cur)
    return "\n\n".join(paragraphs)


__all__ = [
    "MAX_BYTES",
    "MAX_CHARS",
    "MAX_PDF_PAGES",
    "MAX_UNZIPPED_BYTES",
    "SCANNED_MESSAGE",
    "SUPPORTED",
    "DocumentTooLarge",
    "EmptyDocument",
    "IngestError",
    "ScannedPdf",
    "UnsupportedFileType",
    "extract_text",
    "normalise",
    "unwrap",
]
