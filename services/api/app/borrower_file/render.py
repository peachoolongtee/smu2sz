"""Render a BorrowerFile: self-contained printable HTML (bilingual headings, neutral styling), DOCX
(python-docx) and PDF when a local renderer is installed (none is a dependency; see `render_pdf`).

Rendering only formats what `build()` computed: no arithmetic, and every string is escaped.
"""

from __future__ import annotations

import io
from html import escape

from app.borrower_file.models import (
    EVIDENCE_TEXT,
    LENDER_NOTE_ZH,
    STATUS_NOTE,
    STATUS_NOTE_ZH,
    BorrowerFile,
    Evidence,
    Section,
)

_BADGE = {Evidence.OBSERVED: "obs", Evidence.PROJECTED: "proj", Evidence.OWNER_ESTIMATE: "est"}

_CSS = """
:root{--ink:#1d2430;--muted:#5b6472;--line:#d7dce3;--soft:#f4f6f9;--obs:#1f5f8b;--proj:#6a4c93;--est:#9a5b13}
*{box-sizing:border-box}
body{margin:0;background:#fff;color:var(--ink);font:13px/1.45 -apple-system,"Segoe UI",Roboto,"Noto Sans SC",
"PingFang SC","Microsoft YaHei",sans-serif}
main{max-width:960px;margin:0 auto;padding:28px 20px 40px}
header{border-bottom:2px solid var(--ink);padding-bottom:12px;margin-bottom:18px}
h1{font-size:20px;margin:0 0 2px}
.zh{color:var(--muted);font-weight:500}
.meta{display:flex;flex-wrap:wrap;gap:6px 18px;margin-top:10px;color:var(--muted);font-size:12px}
.notice{margin-top:10px;padding:8px 10px;background:var(--soft);border-left:3px solid var(--ink);font-size:12px}
section{margin:22px 0;break-inside:avoid-page}
h2{font-size:15px;margin:0 0 8px;padding-bottom:4px;border-bottom:1px solid var(--line)}
h2 .zh{font-size:13px;margin-left:6px}
table{width:100%;border-collapse:collapse;font-size:12px}
th,td{text-align:left;padding:5px 6px;border-bottom:1px solid var(--line);vertical-align:top}
th{background:var(--soft);font-weight:600}
td.num{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
.fig{display:inline-block;margin:0 10px 2px 0}
.badge{display:inline-block;font-size:9px;letter-spacing:.04em;text-transform:uppercase;padding:0 4px;
border:1px solid currentColor;border-radius:3px;margin-left:4px;vertical-align:1px}
.obs{color:var(--obs)}.proj{color:var(--proj)}.est{color:var(--est)}
.note{color:var(--muted);font-size:11px;margin:6px 0 0}
.rownote{display:block;color:var(--muted);font-size:11px}
footer{margin-top:28px;padding-top:10px;border-top:2px solid var(--ink);font-size:11px;color:var(--muted)}
.hash{font-family:ui-monospace,Menlo,Consolas,monospace;word-break:break-all}
@media print{main{padding:0;max-width:none}body{font-size:11px}section{margin:14px 0}.notice{border-color:#000}}
"""


def _badge(ev: Evidence) -> str:
    cls = _BADGE[ev]
    return f'<span class="badge {cls}" title="{escape(EVIDENCE_TEXT[ev][2])}">{escape(EVIDENCE_TEXT[ev][0])}</span>'


def _fig_html(text: str, ev: Evidence, extra: str = "") -> str:
    cls = " hash" if (extra or "hash " in text) else ""
    return f'<span class="fig{cls}">{escape(text)}{_badge(ev)}</span>'


def _section_html(s: Section) -> str:
    out = [f'<section id="{escape(s.key)}"><h2>{escape(s.title_en)}<span class="zh">{escape(s.title_zh)}</span></h2>']
    if s.columns:
        out.append("<table><thead><tr><th></th>" + "".join(f"<th>{escape(c)}</th>" for c in s.columns)
                   + "</tr></thead><tbody>")
        for r in s.rows:
            cells = "".join(
                f'<td class="num">{escape(f.text)}{_badge(f.evidence)}</td>' for f in r.figures)
            out.append(f"<tr><th scope=\"row\">{escape(r.label)}</th>{cells}</tr>")
        out.append("</tbody></table>")
    else:
        out.append("<table><tbody>")
        for r in s.rows:
            note = f'<span class="rownote">{escape(r.note)}</span>' if r.note else ""
            figs = "".join(_fig_html(f.text, f.evidence) for f in r.figures)
            out.append(f'<tr><th scope="row" style="width:34%">{escape(r.label)}{note}</th><td>{figs}</td></tr>')
        out.append("</tbody></table>")
    out.extend(f'<p class="note">{escape(n)}</p>' for n in s.notes)
    out.append("</section>")
    return "".join(out)


def render_html(bf: BorrowerFile) -> str:
    body = [s for s in bf.sections if s.key != "provenance"]
    prov = next((s for s in bf.sections if s.key == "provenance"), None)
    p = bf.provenance
    head = (
        f"<header><h1>{escape(bf.title_en)}</h1><div class=\"zh\">{escape(bf.title_zh)}</div>"
        f"<div class=\"meta\"><span>File {escape(bf.file_id)}</span><span>Status: prepared, not sent / 已生成，未发送"
        f"</span><span>Organisation {escape(p.org_id)}</span><span>Generated day {bf.generated_day}</span>"
        f"<span>Lender: mocked / 贷款机构：模拟</span></div>"
        f"<div class=\"notice\">{escape(bf.lender_note)} {escape(LENDER_NOTE_ZH)}<br>{escape(STATUS_NOTE)} "
        f"{escape(STATUS_NOTE_ZH)}</div></header>"
    )
    footer = ""
    if prov is not None:
        items = "".join(f"<div>{escape(r.label)}: " + "".join(_fig_html(f.text, f.evidence, "hash") for f in r.figures)
                        + "</div>" for r in prov.rows)
        footer = (f"<footer><strong>{escape(prov.title_en)} / {escape(prov.title_zh)}</strong>{items}"
                  f"<div>Built from verified state only: ledger state, committed facilities, the operations log "
                  f"and commitments. 仅依据已核实的账本状态生成。</div></footer>")
    return (
        "<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\">"
        "<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">"
        f"<title>{escape(bf.title_en)} {escape(bf.file_id)}</title><style>{_CSS}</style></head>"
        f"<body><main>{head}{''.join(_section_html(s) for s in body)}{footer}</main></body></html>"
    )


def render_docx(bf: BorrowerFile) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(9)
    doc.add_heading(bf.title_en, level=0)
    doc.add_paragraph(bf.title_zh)
    doc.add_paragraph(f"File {bf.file_id} · status PREPARED_NOT_SENT · generated day {bf.generated_day} · "
                      f"lender mocked")
    doc.add_paragraph(f"{bf.lender_note} {LENDER_NOTE_ZH}")
    doc.add_paragraph(f"{STATUS_NOTE} {STATUS_NOTE_ZH}")
    for s in bf.sections:
        doc.add_heading(f"{s.title_en} / {s.title_zh}", level=1)
        if s.columns:
            table = doc.add_table(rows=1, cols=len(s.columns) + 1)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = ""
            for i, c in enumerate(s.columns, start=1):
                hdr[i].text = c
            for r in s.rows:
                cells = table.add_row().cells
                cells[0].text = r.label
                for i, f in enumerate(r.figures, start=1):
                    cells[i].text = f"{f.text} [{f.evidence.value}]"
        else:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = "Item", "Value", "Evidence"
            for r in s.rows:
                cells = table.add_row().cells
                cells[0].text = r.label + (f" ({r.note})" if r.note else "")
                cells[1].text = r.value
                cells[2].text = r.evidence
        for n in s.notes:
            doc.add_paragraph(n)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pdf_renderer() -> str | None:
    """The local PDF renderer, if one is installed (no network installs are attempted)."""
    for name in ("weasyprint",):
        try:
            __import__(name)
        except ImportError:
            continue
        return name
    return None


def render_pdf(bf: BorrowerFile) -> bytes | None:
    """PDF via WeasyPrint from the same HTML, or None when it is not installed (it is not a project
    dependency, so PDF is skipped by default; print the HTML to PDF from a browser instead)."""
    if pdf_renderer() != "weasyprint":
        return None
    from weasyprint import HTML  # type: ignore[import-not-found]

    return HTML(string=render_html(bf)).write_pdf()
